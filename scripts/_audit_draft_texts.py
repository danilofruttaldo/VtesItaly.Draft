"""One-shot audit: compare data/DraftTexts.xlsx with cards.json draft_text."""

import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data" / "DraftTexts.xlsx"
CARDS = ROOT / "data" / "cards.json"


def norm_name(s: str) -> str:
    s = (s or "").strip().lower()
    s = re.sub(r"[^a-z0-9]+", "", s)
    return s


def norm_text(s: str) -> str:
    s = (s or "").strip()
    s = re.sub(r"^DRAFT:\s*", "", s, flags=re.I)
    s = s.replace("’", "'")
    s = re.sub(r"\s+", " ", s)
    s = s.strip(" .;,")
    return s.lower()


def main() -> int:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    rows: list[tuple[str, str, str]] = []  # (sheet, name, text)
    for sn in wb.sheetnames:
        ws = wb[sn]
        for r in ws.iter_rows(min_row=2, values_only=True):
            name, text = (r[0] or "").strip() if r[0] else "", (r[1] or "").strip() if r[1] else ""
            if name:
                rows.append((sn, name, text))
    print(f"xlsx rows: {len(rows)}  (across {len(wb.sheetnames)} sheets)")

    cards = json.loads(CARDS.read_text(encoding="utf-8"))
    lib_by_norm: dict[str, dict] = {}
    for c in cards["library"]:
        lib_by_norm[norm_name(c["name"])] = c
    crypt_by_norm: dict[str, dict] = {}
    for c in cards["crypt"]:
        crypt_by_norm[norm_name(c["name"])] = c

    missing_in_cube: list[tuple[str, str]] = []
    missing_text_on_site: list[tuple[str, str, str]] = []  # in cube + draft true but text empty
    not_flagged_draft: list[tuple[str, str]] = []  # in cube but draft=false
    mismatched: list[tuple[str, str, str, str]] = []  # name, sheet, xlsx_text, site_text
    aligned = 0

    seen_xlsx_names = set()
    for sheet, name, xtext in rows:
        key = norm_name(name)
        seen_xlsx_names.add(key)
        if key not in lib_by_norm:
            # might be a crypt entry
            if key in crypt_by_norm:
                missing_in_cube.append((name, f"{sheet} (crypt)"))
            else:
                missing_in_cube.append((name, sheet))
            continue
        c = lib_by_norm[key]
        if not c.get("draft"):
            not_flagged_draft.append((name, sheet))
            continue
        site_text = c.get("draft_text") or ""
        if not site_text:
            missing_text_on_site.append((name, sheet, xtext))
            continue
        if norm_text(site_text) == norm_text(xtext):
            aligned += 1
        else:
            mismatched.append((name, sheet, xtext, site_text))

    # Cards flagged draft=true on site but with no entry in xlsx
    site_draft_names = {norm_name(c["name"]) for c in cards["library"] if c.get("draft")}
    not_in_xlsx = sorted(site_draft_names - seen_xlsx_names)

    print("\n=== SUMMARY ===")
    print(f"xlsx total rows: {len(rows)}")
    print(f"site library cards flagged draft: {len(site_draft_names)}")
    print(f"aligned (xlsx == site): {aligned}")
    print(f"mismatched (xlsx != site): {len(mismatched)}")
    print(f"missing draft_text on site (cube card has draft=true, text empty): {len(missing_text_on_site)}")
    print(f"xlsx entries with no matching cube card: {len(missing_in_cube)}")
    print(f"xlsx entries for cube cards NOT flagged draft on site: {len(not_flagged_draft)}")
    print(f"site DRAFT cards NOT in xlsx: {len(not_in_xlsx)}")

    if missing_text_on_site:
        print("\n--- MISSING TEXT on site (need to import from xlsx) ---")
        for name, sheet, xtext in missing_text_on_site:
            print(f"  [{sheet}] {name} :: {xtext!r}")

    if not_flagged_draft:
        print("\n--- IN CUBE but draft=false (xlsx has DRAFT clause for them) ---")
        for name, sheet in not_flagged_draft:
            print(f"  [{sheet}] {name}")

    if mismatched:
        print("\n--- MISMATCHED text (xlsx vs site) ---")
        for name, sheet, xtext, stext in mismatched:
            print(f"  [{sheet}] {name}")
            print(f"      xlsx: {xtext}")
            print(f"      site: {stext}")

    if not_in_xlsx:
        print("\n--- DRAFT cards on site but NOT in xlsx ---")
        for k in not_in_xlsx:
            # find display name
            for c in cards["library"]:
                if norm_name(c["name"]) == k:
                    print(f"  {c['name']}  --  site text: {c.get('draft_text')!r}")
                    break

    if missing_in_cube:
        print(f"\n--- xlsx entries NOT in current cube ({len(missing_in_cube)}) ---")
        for name, sheet in missing_in_cube[:50]:
            print(f"  [{sheet}] {name}")
        if len(missing_in_cube) > 50:
            print(f"  ... and {len(missing_in_cube) - 50} more")

    return 0


if __name__ == "__main__":
    sys.exit(main())
