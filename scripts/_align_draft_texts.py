"""Align draft_overrides.json + cards.json to data/DraftTexts.xlsx (authoritative).

One-shot helper: reads the official DRAFT texts from DraftTexts.xlsx, updates
draft_overrides.json (canonical key = cube card name) and patches cards.json
draft / draft_text fields directly so the change is visible without rebuilding
from Draft Cube.xlsx.
"""

import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data" / "DraftTexts.xlsx"
CARDS = ROOT / "data" / "cards.json"
OVERRIDES = ROOT / "data" / "draft_overrides.json"

# Sheets ordered older -> newer. Later wins on conflicts of equal length.
SHEET_PRIORITY = [
    "3rd edition",
    "Sword of Caine",
    "Lords of the Night",
    "Twilight Rebellion",
    "Ebony Kingdom",
    "Heirs to the Blood",
    "Keepers of Tradition",
]


def norm_name(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", (s or "").lower())


def clean_xlsx_text(s: str) -> str:
    s = (s or "").strip()
    s = s.replace("{", "[")  # fix '{pro]' typo
    s = re.sub(r"\s+", " ", s)
    return s


def main() -> int:
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    # Build {normalized_name: (text, sheet)} preferring later sheet / longer text
    xlsx_map: dict[str, tuple[str, str]] = {}
    for sn in SHEET_PRIORITY:
        if sn not in wb.sheetnames:
            continue
        for r in wb[sn].iter_rows(min_row=2, values_only=True):
            name, text = r[0], r[1]
            if not name or not text:
                continue
            key = norm_name(str(name))
            cleaned = clean_xlsx_text(str(text))
            prev = xlsx_map.get(key)
            if prev is None or len(cleaned) > len(prev[0]):
                xlsx_map[key] = (cleaned, sn)
            elif len(cleaned) == len(prev[0]):
                # later sheet wins (we iterate in priority order)
                xlsx_map[key] = (cleaned, sn)

    cards = json.loads(CARDS.read_text(encoding="utf-8"))
    library = cards["library"]

    # Map normalized name -> library card (for cube lookup)
    lib_by_norm = {norm_name(c["name"]): c for c in library}

    # Load existing overrides; preserve _comment + entries not in xlsx
    raw_overrides = json.loads(OVERRIDES.read_text(encoding="utf-8"))
    new_overrides: dict[str, str] = {}
    for k, v in raw_overrides.items():
        if k.startswith("_") or not isinstance(v, str):
            new_overrides[k] = v
            continue
        if norm_name(k) not in xlsx_map and norm_name(k) in lib_by_norm:
            # not in xlsx but valid override — keep as-is
            new_overrides[k] = v

    # Apply xlsx truth
    updated = 0
    flagged_new = 0
    for key, (text, _sn) in xlsx_map.items():
        card = lib_by_norm.get(key)
        if not card:
            continue
        canonical = card["name"]
        new_overrides[canonical] = f"DRAFT: {text}" if not text.upper().startswith("DRAFT") else text
        # patch cards.json in-place
        new_text = new_overrides[canonical]
        if card.get("draft_text") != new_text or not card.get("draft"):
            if not card.get("draft"):
                flagged_new += 1
            card["draft"] = True
            card["draft_text"] = new_text
            updated += 1

    # Persist (sorted, _comment first)
    ordered: dict[str, str] = {}
    if "_comment" in new_overrides:
        ordered["_comment"] = new_overrides.pop("_comment")
    for k in sorted(new_overrides, key=str.lower):
        ordered[k] = new_overrides[k]

    OVERRIDES.write_text(
        json.dumps(ordered, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    library.sort(key=lambda c: c["name"].lower())
    CARDS.write_text(
        json.dumps(cards, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    print(f"xlsx unique cards mapped: {len(xlsx_map)}")
    print(f"cards.json entries updated: {updated}")
    print(f"  of which newly flagged draft=true: {flagged_new}")
    print(f"draft_overrides.json entries: {len(ordered) - 1}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
